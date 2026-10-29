"""Load a DOL LCA disclosure file (.xlsx or .csv) into `lca_cases`.

Run:  python -m src.visa.lca_loader --file data/lca/<file> --fiscal-year 2026 [--dry-run]
      python -m src.visa.lca_loader --rekey   (recompute employer_key after alias edits)

Judgment calls (see docs/data_dictionary.md, `lca_cases`):

- **Explicit column allowlist.** The files also carry real people's
  names, emails and phone numbers (EMPLOYER_POC_*, AGENT_ATTORNEY_*,
  LAWFIRM_*, PREPARER_*). Rows are streamed and every cell outside
  `COLUMN_MAP` is dropped the moment its row is read, so those values
  are never held in a DataFrame and never reach the database.
- **Header aliases, matched exactly.** A column DOL has spelled two ways
  (`H-1B_DEPENDENT` / `H_1B_DEPENDENT`) lists both in `COLUMN_MAP`. No
  fuzzy or case-insensitive matching. A file carrying two different
  spellings of one column is refused rather than guessed at.
- **Upsert on case_number, latest load wins.** DOL's quarterly files
  are cumulative within a fiscal year (a Q3 file also holds Q1+Q2), so
  loading Q2 and then Q3 converges on one row per case instead of
  double counting. Load files in chronological order.
- **Defensive parsing.** A bad flag, date, number or wage becomes NULL
  and is counted per column; it never aborts the load. A row with no
  CASE_NUMBER is skipped and counted.
- `employer_key` = `normalize_company_name(EMPLOYER_NAME)` -- the same
  grouping key as `companies.normalized_name`, so employer-level joins
  line up. Names are NOT deduplicated beyond that ("Amazon.com Services
  LLC" and "Amazon Web Services" stay separate employers).
"""

from __future__ import annotations

import argparse
import csv
import logging
import re
from collections import Counter
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert as pg_insert

from src.cleaning.company_normalizer import normalize_company_name
from src.models import LcaCase
from src.utils.db import session_scope

logger = logging.getLogger(__name__)

# Source column (exact DOL header) -> lca_cases column. NOTHING outside
# this map is ever kept -- it is the personal-data firewall. A column may
# accept more than one header spelling: each is listed, matched exactly.
COLUMN_MAP: dict[str, str] = {
    "CASE_NUMBER": "case_number",
    "CASE_STATUS": "case_status",
    "RECEIVED_DATE": "received_date",
    "DECISION_DATE": "decision_date",
    "VISA_CLASS": "visa_class",
    "JOB_TITLE": "job_title",
    "SOC_CODE": "soc_code",
    "SOC_TITLE": "soc_title",
    "FULL_TIME_POSITION": "full_time_position",
    "BEGIN_DATE": "begin_date",
    "TOTAL_WORKER_POSITIONS": "total_worker_positions",
    "NEW_EMPLOYMENT": "new_employment",
    "CONTINUED_EMPLOYMENT": "continued_employment",
    "CHANGE_EMPLOYER": "change_employer",
    "EMPLOYER_NAME": "employer_name",
    "TRADE_NAME_DBA": "trade_name_dba",
    "EMPLOYER_CITY": "employer_city",
    "EMPLOYER_STATE": "employer_state",
    "EMPLOYER_FEIN": "employer_fein",
    "NAICS_CODE": "naics_code",
    "WORKSITE_CITY": "worksite_city",
    "WORKSITE_STATE": "worksite_state",
    "WAGE_RATE_OF_PAY_FROM": "wage_rate_of_pay_from",
    "WAGE_RATE_OF_PAY_TO": "wage_rate_of_pay_to",
    "WAGE_UNIT_OF_PAY": "wage_unit_of_pay",
    "PREVAILING_WAGE": "prevailing_wage",
    "PW_UNIT_OF_PAY": "pw_unit_of_pay",
    "PW_WAGE_LEVEL": "pw_wage_level",
    "H-1B_DEPENDENT": "h1b_dependent",
    "H_1B_DEPENDENT": "h1b_dependent",  # the FY2024-FY2026 files' spelling
    "WILLFUL_VIOLATOR": "willful_violator",
}
# lca_cases column -> every header accepted for it, in COLUMN_MAP order.
ALIASES: dict[str, tuple[str, ...]] = {}
for _source, _column in COLUMN_MAP.items():
    ALIASES[_column] = ALIASES.get(_column, ()) + (_source,)
COLUMNS = tuple(ALIASES)

# Without these a row can't be keyed or attributed -- fail fast if absent.
REQUIRED_COLUMNS = ("case_number", "case_status", "employer_name")

_FLAG_COLUMNS = {"full_time_position", "h1b_dependent", "willful_violator"}
_DATE_COLUMNS = {"received_date", "decision_date", "begin_date"}
_INT_COLUMNS = {
    "total_worker_positions",
    "new_employment",
    "continued_employment",
    "change_employer",
}
_WAGE_COLUMNS = {"wage_rate_of_pay_from", "wage_rate_of_pay_to", "prevailing_wage"}

_TRUE = {"y", "yes", "true", "t", "1"}
_FALSE = {"n", "no", "false", "f", "0"}
_DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%m/%d/%y")
_WAGE_JUNK = re.compile(r"[$,\s]")

BATCH_SIZE = 1000


class MissingColumnsError(ValueError):
    """The file lacks a column the loader can't work without."""


class AmbiguousColumnsError(ValueError):
    """The file carries two different header spellings of one column."""


def describe(column: str) -> str:
    """A column and every header tried for it: `h1b_dependent (H-1B_DEPENDENT, H_1B_DEPENDENT)`."""
    return f"{column} ({', '.join(ALIASES[column])})"


@dataclass
class LoadResult:
    rows_read: int = 0
    rows_loaded: int = 0  # distinct case_numbers upserted from this file
    skipped_no_case_number: int = 0
    duplicate_case_numbers_in_file: int = 0
    missing_columns: list[str] = field(default_factory=list)  # lca_cases columns
    bad_cells: Counter = field(default_factory=Counter)  # column -> count nulled
    by_case_status: Counter = field(default_factory=Counter)
    by_visa_class: Counter = field(default_factory=Counter)


# --------------------------------------------------------------------- #
# cell parsers -- each returns (value, ok). ok=False means "bad cell".


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def parse_text(value: Any) -> tuple[str | None, bool]:
    if _blank(value):
        return None, True
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel turns NAICS 541511 into 541511.0
    return str(value).strip(), True


def parse_flag(value: Any) -> tuple[bool | None, bool]:
    if _blank(value):
        return None, True
    if isinstance(value, bool):
        return value, True
    s = str(value).strip().lower()
    if s in _TRUE:
        return True, True
    if s in _FALSE:
        return False, True
    return None, False


def parse_date(value: Any) -> tuple[date | None, bool]:
    if _blank(value):
        return None, True
    if isinstance(value, datetime):
        return value.date(), True
    if isinstance(value, date):
        return value, True
    s = str(value).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date(), True
        except ValueError:
            continue
    return None, False


def parse_int(value: Any) -> tuple[int | None, bool]:
    if _blank(value):
        return None, True
    try:
        d = Decimal(_WAGE_JUNK.sub("", str(value)))
    except InvalidOperation:
        return None, False
    if d != d.to_integral_value() or not d.is_finite():
        return None, False
    return int(d), True


def parse_wage(value: Any) -> tuple[Decimal | None, bool]:
    """'$85,000.00' / 85000 / '85000' -> Decimal. Ranges or words -> bad."""
    if _blank(value):
        return None, True
    try:
        d = Decimal(_WAGE_JUNK.sub("", str(value)))
    except InvalidOperation:
        return None, False
    if not d.is_finite() or d < 0:
        return None, False
    return d.quantize(Decimal("0.01")), True


def _parser_for(column: str):
    if column in _FLAG_COLUMNS:
        return parse_flag
    if column in _DATE_COLUMNS:
        return parse_date
    if column in _INT_COLUMNS:
        return parse_int
    if column in _WAGE_COLUMNS:
        return parse_wage
    return parse_text


# --------------------------------------------------------------------- #
# readers -- yield {lca_cases column: raw value} for allowlisted columns only


def _allowlisted_indexes(header: list[Any], source: str) -> dict[int, str]:
    """Header position -> lca_cases column, for exact allowlisted headers only."""
    idx: dict[int, str] = {}
    spelled: dict[str, str] = {}  # column -> the header that supplied it
    for i, name in enumerate(header):
        name = (str(name).strip() if name is not None else "").lstrip("﻿")
        column = COLUMN_MAP.get(name)
        if column is None:
            continue
        if column in spelled and spelled[column] != name:
            raise AmbiguousColumnsError(
                f"{source} has both {spelled[column]!r} and {name!r}, two spellings of "
                f"{column}: keep one of them and reload"
            )
        spelled[column] = name
        idx[i] = column
    return idx


def _iter_xlsx(path: Path) -> Iterator[tuple[list[str], dict[str, Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        header = next(rows, None) or []
        idx = _allowlisted_indexes(list(header), path.name)
        present = list(idx.values())
        for row in rows:
            yield present, {col: row[i] for i, col in idx.items() if i < len(row)}
    finally:
        wb.close()


def _iter_csv(path: Path) -> Iterator[tuple[list[str], dict[str, Any]]]:
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as fh:
        reader = csv.reader(fh)
        header = next(reader, None) or []
        idx = _allowlisted_indexes(header, path.name)
        present = list(idx.values())
        for row in reader:
            yield present, {col: row[i] for i, col in idx.items() if i < len(row)}


def iter_allowlisted_rows(path: Path) -> Iterator[tuple[list[str], dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _iter_xlsx(path)
    if suffix == ".csv":
        return _iter_csv(path)
    raise ValueError(f"unsupported file type {suffix!r}: expected .xlsx or .csv")


# --------------------------------------------------------------------- #


def parse_row(raw: dict[str, Any], bad_cells: Counter) -> dict[str, Any]:
    """`raw` is keyed by lca_cases column (the readers translate headers)."""
    out: dict[str, Any] = {}
    for column in COLUMNS:
        value, ok = _parser_for(column)(raw.get(column))
        if not ok:
            bad_cells[column] += 1
        out[column] = value
    return out


def load_file(
    path: Path,
    fiscal_year: int,
    *,
    dry_run: bool = False,
    database_url: str | None = None,
) -> LoadResult:
    """Parse `path` and upsert its cases. Returns counts; never raises on bad cells."""
    path = Path(path)
    result = LoadResult()
    records: dict[str, dict[str, Any]] = {}  # case_number -> row; last occurrence wins
    loaded_at = datetime.now(UTC)
    checked_header = False

    for present, raw in iter_allowlisted_rows(path):
        if not checked_header:
            missing = [c for c in COLUMNS if c not in present]
            hard = [c for c in REQUIRED_COLUMNS if c in missing]
            if hard:
                raise MissingColumnsError(
                    f"{path.name} lacks required column(s), headers tried: "
                    + "; ".join(describe(c) for c in hard)
                )
            result.missing_columns = missing
            if missing:
                logger.warning(
                    "%s lacks allowlisted column(s), loaded as NULL (headers tried): %s",
                    path.name,
                    "; ".join(describe(c) for c in missing),
                )
            checked_header = True

        if all(_blank(v) for v in raw.values()):
            continue  # trailing empty spreadsheet rows
        result.rows_read += 1
        row = parse_row(raw, result.bad_cells)
        if not row["case_number"]:
            result.skipped_no_case_number += 1
            continue
        if row["case_number"] in records:
            result.duplicate_case_numbers_in_file += 1
        row["employer_key"] = (
            normalize_company_name(row["employer_name"]) if row["employer_name"] else None
        )
        row["fiscal_year"] = fiscal_year
        row["source_file"] = path.name
        row["loaded_at"] = loaded_at
        records[row["case_number"]] = row

    for row in records.values():
        result.by_case_status[row["case_status"] or "(none)"] += 1
        result.by_visa_class[row["visa_class"] or "(none)"] += 1
    result.rows_loaded = len(records)

    if not dry_run and records:
        _upsert(list(records.values()), database_url)
    return result


def _upsert(rows: list[dict[str, Any]], database_url: str | None) -> None:
    update_cols = [c for c in rows[0] if c != "case_number"]
    with session_scope(database_url) as session:
        for start in range(0, len(rows), BATCH_SIZE):
            stmt = pg_insert(LcaCase).values(rows[start : start + BATCH_SIZE])
            stmt = stmt.on_conflict_do_update(
                index_elements=["case_number"],
                set_={c: stmt.excluded[c] for c in update_cols},
            )
            session.execute(stmt)


def rekey_employers(database_url: str | None = None) -> int:
    """Recompute employer_key for every stored case from employer_name.

    Run after promoting pairs into company_aliases.json -- cheaper than
    reloading every file, and keeps the key recomputable. Returns rows changed.
    """
    changed = 0
    with session_scope(database_url) as session:
        for case in session.execute(select(LcaCase)).scalars():
            key = normalize_company_name(case.employer_name) if case.employer_name else None
            if key != case.employer_key:
                case.employer_key = key
                changed += 1
    return changed


def summarize(result: LoadResult, dry_run: bool) -> str:
    lines = [
        f"rows read: {result.rows_read}",
        f"cases {'parsed (dry run, nothing written)' if dry_run else 'upserted'}: {result.rows_loaded}",
        f"skipped (no CASE_NUMBER): {result.skipped_no_case_number}",
        f"duplicate CASE_NUMBERs within file (last kept): {result.duplicate_case_numbers_in_file}",
    ]
    if result.missing_columns:
        lines.append(
            "allowlisted columns absent from file (NULL; headers tried): "
            + "; ".join(describe(c) for c in result.missing_columns)
        )
    lines.append(
        "bad cells nulled: " + (str(dict(result.bad_cells)) if result.bad_cells else "none")
    )
    lines.append("by case_status:")
    lines += [f"  {k:<22} {v}" for k, v in result.by_case_status.most_common()]
    lines.append("by visa_class:")
    lines += [f"  {k:<22} {v}" for k, v in result.by_visa_class.most_common()]
    return "\n".join(lines)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    argp = argparse.ArgumentParser(description="Load a DOL LCA disclosure file into lca_cases.")
    argp.add_argument("--file", type=Path, help=".xlsx or .csv under data/lca/")
    argp.add_argument("--fiscal-year", type=int, help="federal FY of the file")
    argp.add_argument("--dry-run", action="store_true", help="parse and count, write nothing")
    argp.add_argument(
        "--rekey",
        action="store_true",
        help="recompute employer_key for stored cases (after editing company_aliases.json)",
    )
    args = argp.parse_args()

    if args.rekey:
        print(f"employer_key changed on {rekey_employers()} cases")
        return
    if args.file is None or args.fiscal_year is None:
        argp.error("--file and --fiscal-year are required (or use --rekey)")

    result = load_file(args.file, args.fiscal_year, dry_run=args.dry_run)
    print(summarize(result, args.dry_run))


if __name__ == "__main__":
    main()
